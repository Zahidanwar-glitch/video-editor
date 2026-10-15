"""Step C5 - design calculation report (Word + PDF via LibreOffice).

Sections: design basis (codes, conditions, envelope, internal loads, assumptions), room
schedule, room-by-room loads with a worked example, ventilation + exhaust, zoning and
equipment selection, duct sizing and air terminals, piping, checks, clause register.
Every table that applies a code rule carries the clause reference.

CLI: python report.py -w WORK   -> WORK/out/HVAC_Design_Report.docx / .pdf
"""
import argparse
import os
import shutil
import subprocess
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import common as C

RED = RGBColor(0xC0, 0, 0)


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcol)
    tcPr.append(sh)


def table(doc, headers, rows, widths=None, size=7.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(size)
        shade(c, "DDEBF7")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def clause_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("Code reference: " + text)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)


def build(work):
    p = C.project(work)
    rooms = C.load(work, "rooms.json")
    L = C.load(work, "loads.json")
    V = C.load(work, "ventilation.json")
    D = C.load(work, "design.json")
    chk = C.load(work, "checks.json", default={"checks": []})
    A = C.load(work, "assumptions.json", default={"items": []})["items"]
    Q = [q for q in C.load(work, "open_questions.json", default={"items": []})["items"] if not q["answered"]]
    U = C.Units(p["units"])
    note = p["design"]["drawing"]["preliminary_note"]
    dc = L["design_conditions"]
    cl = lambda k: C.clause(p, k)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9.5)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    hp = sec.header.paragraphs[0]
    hr = hp.add_run(f"{p['project']['name']} - HVAC design calculation report - PRELIMINARY")
    hr.font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(note)
    fr.font.size = Pt(7.5)
    fr.font.color.rgb = RED
    fr.bold = True

    # cover
    doc.add_paragraph().add_run("\n\n")
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("HVAC DESIGN CALCULATION REPORT")
    r.bold = True
    r.font.size = Pt(24)
    for line in (p["project"]["name"], f"{p['project']['location']['city']}, {p['project']['location']['country']}",
                 f"Building type: {p['project']['building_type']} - {len(p['project'].get('floors', []))} floor(s)",
                 f"Code basis: {p['codes']['label']}   |   Units: {p['units']}", f"Issue: P0 (preliminary)   {date.today().isoformat()}"):
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run(line).font.size = Pt(13)
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = q.add_run("\n" + note)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RED
    doc.add_page_break()

    # 1 design basis
    doc.add_heading("1. Design basis", 1)
    doc.add_heading("1.1 Codes and standards", 2)
    table(doc, ["Ref", "Code / standard"], [[c["id"], c["title"]] for c in p["codes"]["codes"]], [3, 22])
    if p["codes"].get("local_codes"):
        lc = p["codes"]["local_codes"].get(p["project"]["location"]["country"].lower())
        if lc:
            doc.add_paragraph("Local codes: " + "; ".join(lc))
    doc.add_heading("1.2 Design conditions", 2)
    table(doc, ["Parameter", "Value", "Basis"], [
        ["Outdoor summer design (0.4 %)", f"{U.temp(dc['to_c'])} DB / {U.temp(dc['twb_c'])} MCWB, W = {dc['wo'] * 1000:.1f} g/kg", dc["source"]],
        ["Daily range", f"{dc['dr_k']} K", dc["source"]],
        ["Outdoor winter design (99.6 %)", U.temp(dc["th_c"]), dc["source"]],
        ["Indoor cooling", f"{U.temp(dc['ti_c'])}, {dc['rhi']} % RH (W = {dc['wi'] * 1000:.1f} g/kg)", cl("design_conditions")],
        ["Indoor heating", U.temp(dc["tih_c"]), cl("design_conditions")],
        ["Site elevation / pressure", f"{dc['elev_m']:.0f} m / {dc['p_kpa']} kPa; air factors {dc['cs']:.3f} W/(L/s·K), {dc['cl']:.0f} W/(L/s·kg/kg)", "ASHRAE Fundamentals Ch.1"],
        ["Latitude / design month", f"{dc['lat']}° / July (CLTD tables)", cl("load_method")]], [6, 10, 10])
    clause_note(doc, cl("design_conditions"))
    doc.add_heading("1.3 Building envelope", 2)
    Uv = L["u_values"]
    table(doc, ["Element", "U-value / property", "Basis"], [
        ["External wall", U.u_val(Uv["wall_ext"]) + f" (CLTD wall group {Uv['wall_group']})", L["u_set"]],
        ["Roof", U.u_val(Uv["roof"]) + f" ({Uv['roof_type']})", L["u_set"]],
        ["Glazing", f"{U.u_val(Uv['glazing_u'])}, SHGC {Uv['glazing_shgc']} (SC {Uv['glazing_shgc'] / 0.87:.2f})", L["u_set"]],
        ["Slab edge (heating)", f"F = {Uv['slab_edge_f_w_mk']} W/m·K", L["u_set"]]], [6, 10, 10])
    clause_note(doc, cl("envelope_u"))
    doc.add_heading("1.4 Internal loads", 2)
    stt = C.ref("space_types.json")
    used = sorted({r["space_type"] for r in L["rooms"]})
    rows = []
    for k in used:
        s = stt["types"].get(k, {})
        occ = s.get("occupants", {})
        rows.append([k, f"{occ.get('fixed')} persons" if "fixed" in occ else f"{occ.get('m2_per_person')} m²/person", s.get("activity"),
                     s.get("lighting_w_m2"), f"{s.get('equip_w_m2')} + {s.get('equip_fixed_w')} W", s.get("vent_category"), s.get("exhaust"), "yes" if s.get("conditioned") else "no"])
    table(doc, ["Space type", "Occupancy", "Activity", "Lighting W/m²", "Equipment W/m² + fixed", "OA category", "Exhaust", "Cooled"], rows)
    clause_note(doc, cl("internal_gains"))
    doc.add_heading("1.5 Assumptions", 2)
    doc.add_paragraph("All assumptions are logged in assumptions.md. Items marked 'confirm' must be confirmed by the engineer of record.")
    table(doc, ["ID", "Step", "Assumption", "Source", "Status"], [[a["id"], a["step"], a["text"], a["source"], "CONFIRM" if a["confirm"] else "default"] for a in A], [1.5, 2, 13, 6, 2], size=7)
    if Q:
        doc.add_heading("1.6 Open questions", 2)
        table(doc, ["ID", "Question"], [[q["id"], q["question"]] for q in Q], [2, 22])

    # 2 room schedule
    doc.add_heading("2. Room schedule (Gate 1 approved)", 1)
    rows = []
    for f in rooms["floors"]:
        for r in f.get("rooms", []):
            if f.get("kind") == "roof":
                continue
            rows.append([r["id"], r["name"], r["space_type"], f"{r['area_m2']:.2f}", r.get("height_m"),
                         ", ".join(f"{e['orientation']} {e['length_m']}" for e in r["exterior_walls"]) or "-",
                         ", ".join(f"{w['tag']} {w['orientation']} {w['width_m']}x{w['height_m']}" for w in r.get("windows", [])) or "-",
                         f"{r.get('roof_exposed_frac', 0):.0%}", "yes" if r.get("is_conditioned") else "no"])
    table(doc, ["ID", "Room", "Type", "Area m²", "Ceiling m", "Exterior walls (orient. m)", "Windows", "Roof exp.", "Cooled"], rows, size=7)

    # 3 loads
    doc.add_heading("3. Cooling and heating loads", 1)
    doc.add_paragraph(
        "Method: CLTD/SCL/CLF (ASHRAE Fundamentals 1997 Ch.28). Conduction Q = U·A·CLTDc with CLTDc = (CLTD + LM)·K + (25.5 − Ti) + (Tm − 29.4), "
        "Tm = To − DR/2. Glass solar Q = A·SC·SHGF·CLF. People/lights/equipment with CLF = 1.0. Infiltration Qs = Cs·V·ΔT, Ql = Cl·V·ΔW. "
        "Loads are computed for solar hours 8-20; each room is sized at its own peak, zones and plant at the coincident (block) peak. "
        "Heating: Q = ΣU·A·(Ti − To,w) + infiltration + slab edge, no internal gains.")
    clause_note(doc, cl("load_method") + "; " + cl("infiltration"))
    rows = []
    for r in L["rooms"]:
        if not r["conditioned"]:
            continue
        c = r["components_w"]
        rows.append([r["id"], r["name"], r["area_m2"], c["wall"], c["roof"], c["glass_cond"], c["glass_solar"], c["people_s"], c["lights"], c["equip_s"], c["infil_s"],
                     r["sensible_w"], r["latent_w"], r["total_w"], r["w_per_m2"], f"{r['peak_hour']}:00", U.flow(r["supply_ls"]), r["heating_w"]])
    table(doc, ["ID", "Room", "m²", "Walls", "Roof", "Glass cond.", "Glass solar", "People S", "Lights", "Equip S", "Infil S", "Sensible W", "Latent W", "Total W", "W/m²", "Peak", "Supply", "Heating W"], rows, size=6.5)
    s = L["summary"]
    doc.add_paragraph(f"Conditioned area {s['conditioned_area_m2']} m²; sum of room peaks {U.power(s['sum_room_peaks_kw'])}; block peak {U.power(s['block_peak_kw'])} at {s['block_peak_hour']}:00 "
                      f"({s['block_peak_kw'] * 1000 / max(s['conditioned_area_m2'], 1):.0f} W/m²); heating {U.power(s['heating_kw'])}. Safety factors are applied at equipment selection (section 5).")
    big = max((r for r in L["rooms"] if r["conditioned"]), key=lambda r: r["total_w"])
    doc.add_heading(f"3.1 Worked example - {big['id']} {big['name']} at {big['peak_hour']}:00", 2)
    rows = [[f"Wall {w['orientation']}", f"U {w['u']} × A {w['net_m2']} m² (gross {w['gross_m2']})"] for w in big["envelope"]["walls"]]
    rows += [[f"Window {w['tag']} {w['orientation']}", f"cond. U {w['u']} × {w['area_m2']} m²; solar {w['area_m2']} × SC {w['sc']} × SHGF {w['shgf']} × CLF"] for w in big["envelope"]["windows"]]
    rows += [["Roof", f"U {L['u_values']['roof']} × {big['roof_area_m2']} m²"], ["People", f"{big['people']} persons"], ["Infiltration", f"{big['infil_ach']} ACH × {big['volume_m3']} m³ = {big['infil_ls']} L/s"]]
    rows += [[k, f"{v} W"] for k, v in big["components_w"].items()]
    rows += [["Room sensible / latent / total", f"{big['sensible_w']} / {big['latent_w']} / {big['total_w']} W"],
             ["Supply airflow", f"{big['peak_sensible_w']} W / ({dc['cs']:.3f} × ({dc['ti_c']} − {p['design']['loads']['supply_air_temp_c']})) = {U.flow(big['supply_ls'], 1)}"]]
    table(doc, ["Item", "Calculation"], rows, [7, 18])

    # 4 ventilation
    doc.add_heading("4. Ventilation and exhaust", 1)
    doc.add_paragraph("Ventilation rate procedure: Vbz = Rp·Pz + Ra·Az; Voz = Vbz / Ez; single-zone Vot = Voz; multiple-zone Vot = (D·ΣRp·Pz + ΣRa·Az)/Ev.")
    clause_note(doc, V["basis"]["clause_oa"] + "; " + V["basis"]["clause_ez"] + "; " + V["basis"]["clause_ev"])
    rows = []
    for zid, z in V["zones"].items():
        for r in z["rooms"]:
            rows.append([zid, r["room"], r["category"], r["people"], r["area_m2"], r["Rp_Pz"], r["Ra_Az"], r["Vbz"], z["Ez"], z["Ev"], z["Vot_ls"], z["oa_method"]])
    table(doc, ["Zone", "Room", "Category", "Pz", "Az m²", "Rp·Pz", "Ra·Az", "Vbz L/s", "Ez", "Ev", "Zone Vot L/s", "OA method"], rows, size=7)
    table(doc, ["Room", "Name", "Rule", "Exhaust L/s", "Mode", "Clause"], [[k, e["room"], e["rule"], e["ls"], e["mode"], e["clause"]] for k, e in V["exhaust"].items()], size=7)
    t = V["totals"]
    doc.add_paragraph(f"Building outdoor air {U.flow(t['oa_ls'])}; exhaust {U.flow(t['exhaust_ls'])} (net {t['balance_ls']:+.0f} L/s)."
                      + (f" ASHRAE 62.2 whole-house cross-check: {U.flow(t['res_62_2_check_ls'])} - design OA exceeds it." if t.get("res_62_2_check_ls") else ""))

    # 5 equipment
    doc.add_heading("5. Zoning and equipment selection", 1)
    doc.add_paragraph(f"Required capacity = coincident zone load (rooms + outdoor air) × (1 + SF), SF cooling {D['sf_cooling']:.0%}, heating {D['sf_heating']:.0%}. "
                      f"Nominal capacities × {D['derate']} for high-ambient (T3) operation. Selection = smallest catalogue unit meeting total, sensible and airflow.")
    clause_note(doc, cl("safety_factor") + "; " + cl("equipment_efficiency"))
    rows = []
    for z in D["zones"]:
        u = z["unit"]
        rows.append([z["id"], ", ".join(z["rooms"]), z["system"], u["peak_hour"], u["room_kw"], u["oa_kw"], u["block_kw"], f"{D['sf_cooling']:.0%}", u["req_total_kw"], u["req_sensible_kw"],
                     U.flow(u["req_airflow_ls"]), f"{u['tag']} {u['model']}", u["derated_kw"], u["sensible_kw"], U.flow(u["airflow_ls"]), f"{u['oa_frac']:.0%}", u["power_kw"], f"{u['heat_kw']} / {u['heat_req_kw']}"])
    table(doc, ["Zone", "Rooms", "System", "Peak h", "Rooms kW", "OA kW", "Block kW", "SF", "Req. kW", "Req. sens kW", "Req. air", "Selected", "Derated kW", "Sens kW", "Unit air", "OA %", "Input kW", "Heat kW sel/req"], rows, size=6.5)
    table(doc, ["Tag", "Model", "System", "Serves", "Nominal", "Input kW", "Location"], [[o["tag"], o["model"], o["system"], ", ".join(o["serves"]), U.power(o["cool_kw"]), o["power_kw"], o.get("floor", "TBC")] for o in D["outdoor_units"]])

    # 6 ducts
    doc.add_heading("6. Duct sizing and air terminals", 1)
    ds = C.ref("duct_sizing.json")
    cls = "residential" if C.building_class(p) == "residential" else "commercial"
    doc.add_paragraph(f"Equal friction method, target {U.fric(ds['friction_target_pa_m'][cls])}, velocity limits (m/s) {ds['velocity_limits_m_s'][cls]}. Darcy-Weisbach with Altshul-Tsal friction factor, "
                      f"roughness {ds['air']['roughness_mm']} mm; rectangular sizes from De = 1.30(ab)^0.625/(a+b)^0.25; fittings +{p['design']['ducts']['fittings_equivalent_length_factor']:.0%} equivalent length.")
    clause_note(doc, cl("duct_sizing") + "; " + cl("duct_construction"))
    for a in D.get("air_systems", []):
        doc.add_heading(f"{a['zone']} - {a['unit']} ({a['floor']})", 3)
        rows = [[s["id"], s["kind"], s.get("serves") or "-", U.flow(s["flow_ls"]), U.rect(s["w_mm"], s["h_mm"]) if s["shape"] == "rect" else f"Ø{U.size(s['d_mm'])} {s['shape']}",
                 s["de_mm"], U.vel(s["velocity_m_s"]), U.fric(s["friction_pa_m"]), s["length_m"], U.pa(s["dp_pa"], 1), s["gauge"]] for s in a["segments"] + a["return"]["segments"] + (a["fresh_air"]["segments"] if a.get("fresh_air") else [])]
        table(doc, ["Segment", "Type", "Serves", "Flow", "Size", "De mm", "Velocity", "Friction", "Length m", "Δp", "Gauge"], rows, size=6.5)
        doc.add_paragraph(f"Critical path: required ESP {U.pa(a['esp_required_pa'])} vs available {U.pa(a['esp_available_pa'])} - {'OK' if a['esp_ok'] else 'NOT OK'}.")
    rows = [[d["tag"], d["room"], d["model"], U.flow(d["flow_ls"]), d["nc"], d["throw_m"]] for a in D.get("air_systems", []) for d in a["diffusers"]]
    doc.add_heading("6.1 Supply diffuser schedule", 2)
    table(doc, ["Tag", "Room", "Model", "Flow", "NC", "Throw m"], rows, size=7)
    clause_note(doc, cl("noise"))
    doc.add_heading("6.2 Exhaust systems", 2)
    rows = [[e["fan"]["tag"], e["room"], U.flow(e["flow_ls"]), e["fan"]["model"], f"Ø{U.size(e['segments'][0]['d_mm'])}", e["segments"][0]["length_m"], U.vel(e["segments"][0]["velocity_m_s"]),
             U.pa(e["fan"]["esp_req_pa"]), e["discharge"]["type"], e["clause"]] for e in D.get("exhaust_systems", [])]
    table(doc, ["Fan", "Room", "Flow", "Model", "Duct", "Length m", "Velocity", "ESP req.", "Discharge", "Clause"], rows, size=7)

    # 7 pipes
    doc.add_heading("7. Piping", 1)
    P = D.get("pipes", {})
    if P.get("refrigerant"):
        doc.add_heading("7.1 Refrigerant piping", 2)
        table(doc, ["Run", "Unit → ODU", "Capacity", "Liquid / gas OD mm", "Horiz. floor m", "Vertical m", "Roof m", "Total m", "Limit m", "Status"],
              [[r["tag"], f"{r['unit']} → {r['odu']}", U.power(r["capacity_kw"]), f"{r['liquid_mm']} / {r['gas_mm']} ({r['liquid_in']}\" / {r['gas_in']}\")", r["horiz_floor_m"], r["vertical_m"], r["horiz_roof_m"],
                r["total_m"], r.get("max_length_m", "-"), "OK" if r["ok"] else "CHECK"] for r in P["refrigerant"]], size=7)
        clause_note(doc, cl("refrigerant_piping"))
    if P.get("chw"):
        doc.add_heading("7.2 Chilled water piping", 2)
        table(doc, ["Run", "Flow L/s", "DN", "Velocity m/s", "Friction Pa/m", "Length m"], [[c["tag"], c["flow_ls"], c["dn"], c["velocity_m_s"], c["friction_pa_m"], c["length_m"]] for c in P["chw"]])
        clause_note(doc, cl("chw_piping"))
    if P.get("condensate"):
        doc.add_heading("7.3 Condensate drains", 2)
        table(doc, ["Run", "Unit", "To", "Code min mm", "Design mm", "Length m", "Slope %"], [[c["tag"], c["unit"], c["to"], c["code_min_mm"], c["dn_mm"], c["length_m"], c["slope_pct"]] for c in P["condensate"]], size=7)
        clause_note(doc, cl("condensate"))

    # 8 checks, 9 clauses
    doc.add_heading("8. Design checks", 1)
    table(doc, ["Check", "Status", "Detail"], [[c["check"], c["status"], c["detail"]] for c in chk["checks"]], [6, 2, 17])
    doc.add_heading("9. Code clause register", 1)
    table(doc, ["Rule", "Clause applied"], [[k, v] for k, v in p["codes"]["clauses"].items()], [5, 20])
    prev = os.path.join(work, "out", "preview_M-101.png")
    if os.path.exists(prev):
        doc.add_page_break()
        doc.add_heading("Appendix - sample sheet (M-101)", 1)
        doc.add_picture(prev, width=Cm(24))
    os.makedirs(os.path.join(work, "out"), exist_ok=True)
    path = os.path.join(work, "out", "HVAC_Design_Report.docx")
    doc.save(path)
    pdf = to_pdf(path)
    print(f"report: {path}" + (f" + {os.path.basename(pdf)}" if pdf else " (PDF conversion skipped - LibreOffice not found)"))


def to_pdf(path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    out = os.path.dirname(path)
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", out, path], check=True, capture_output=True, timeout=240)
    pdf = os.path.splitext(path)[0] + ".pdf"
    return pdf if os.path.exists(pdf) else None


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-w", "--work", required=True)
    build(ap.parse_args().work)
