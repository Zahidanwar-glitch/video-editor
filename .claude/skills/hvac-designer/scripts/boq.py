"""Step C3 - quantities-only BOQ (no prices) to Excel.

Columns: Item No., Description, Specification, Unit, Quantity, Waste %, Qty incl. allowance
(formula, so the waste % is editable), Floor/Zone, Remarks (tags for traceability).
Sections: equipment; ductwork (m² by gauge); flexible duct; air terminals; dampers; fans;
refrigerant / CHW pipes and fittings; insulation; condensate; supports; controls; T&C.

CLI: python boq.py -w WORK   -> WORK/out/HVAC_BOQ.xlsx, WORK/boq.json
"""
import argparse
import math
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import common as C


def build(work):
    p = C.project(work)
    design = C.load(work, "design.json")
    R = C.ref("boq_rules.json")
    PS = C.ref("pipe_sizing.json")
    DSZ = C.ref("duct_sizing.json")
    U = C.Units(p["units"])
    std = DSZ["gauge_standard_by_code_set"].get(p["project"]["code_set"], "SMACNA")
    W = R["waste_pct"]
    sections = []

    def sec(title, key):
        s = {"title": title, "key": key, "rows": []}
        sections.append(s)
        return s

    def row(s, desc, spec, unit, qty, zone="", tags=(), remarks=""):
        if qty <= 0:
            return
        tg = sorted(set(tags))
        rem = remarks
        if tg:
            rem = (rem + "; " if rem else "") + "Tags: " + ", ".join(tg)
        s["rows"].append({"desc": desc, "spec": spec, "unit": unit, "qty": round(qty, 2), "zone": zone, "remarks": rem, "tags": tg})

    # A. equipment
    s = sec("A. HVAC EQUIPMENT", "equipment")
    for z in design["zones"]:
        u = z["unit"]
        lbl = {"wall_split": "Wall-mounted split indoor unit", "ducted_split": "Ducted split indoor unit (concealed)", "vrf_indoor_ducted": "VRF ducted indoor unit",
               "fcu_ducted": "Chilled-water ducted fan coil unit", "packaged_rooftop": "Packaged rooftop unit"}.get(u["type"], u["type"])
        row(s, lbl, f"{u['model']}: {U.power(u['nominal_kw'])} nominal cooling, {U.flow(u['airflow_ls'])}, ESP {u['esp_pa']} Pa, heat {U.power(u['heat_kw'])}, {p['preferences'].get('refrigerant') or p['design']['equipment']['refrigerant']}",
            "No.", 1, z["id"], [u["tag"]])
        if z.get("fa_fan"):
            f = z["fa_fan"]
            row(s, "Through-wall fresh air fan with G4 filter and weather hood", f"{f['model']}: {U.flow(f['airflow_ls'])}, {f['power_w']} W", "No.", 1, z["id"], [f["tag"]])
    for o in design["outdoor_units"]:
        lbl = "Air-cooled chiller" if o["system"] == "chilled_water" else "Outdoor condensing unit (heat pump)"
        row(s, lbl, f"{o['model']}: {U.power(o['cool_kw'])} nominal, {o['power_kw']} kW input, T3 rated, anti-vibration mounts", "No.", 1, o.get("floor", "Roof"), [o["tag"]])
    if design.get("plant"):
        for pm in design["plant"]["pumps"]:
            row(s, f"Chilled water pump ({pm['duty']})", f"{U.flow(pm['flow_ls'])}, head by detailed design", "No.", 1, "Plant", [pm["tag"]])

    # B. ductwork m2 by gauge, C. flex
    sb = sec(f"B. GALVANISED STEEL DUCTWORK ({std})", "ductwork")
    sc = sec("C. FLEXIBLE DUCT", "flexible_duct")
    duct_area = defaultdict(float)
    flex = defaultdict(float)
    ins_area = defaultdict(float)
    rigid_len = defaultdict(float)
    vcd = defaultdict(list)
    all_segs = []
    for a in design.get("air_systems", []):
        segs = [(x, "supply") for x in a["segments"]] + [(x, "return") for x in a["return"]["segments"]]
        if a.get("fresh_air"):
            segs += [(x, "fresh_air") for x in a["fresh_air"]["segments"]]
        all_segs += [(x, sy, a["floor"]) for x, sy in segs]
        for sname, v in [(s_, a["floor"]) for s_ in a["vcd_branches"]]:
            seg = next(x for x in a["segments"] if x.get("serves") == sname and x["kind"] == "branch")
            vcd[(a["floor"], U.rect(seg["w_mm"], seg["h_mm"]))].append(a["zone"])
    for e in design.get("exhaust_systems", []):
        all_segs += [(x, "exhaust", e["floor"]) for x in e["segments"]]
    for sg, sy, fl in all_segs:
        if sg["shape"] == "flex":
            flex[(fl, sg["d_mm"])] += sg["length_m"]
            continue
        shape = "rectangular" if sg["shape"] == "rect" else "round spiral"
        duct_area[(fl, shape, sg["gauge"])] += sg["area_m2"]
        rigid_len[fl] += sg["length_m"]
        if R["duct_insulation"].get(sy):
            ins_area[(fl, R["duct_insulation"][sy])] += sg["area_m2"]
    for (fl, shape, g), a_ in sorted(duct_area.items()):
        row(sb, f"Galvanised sheet steel {shape} duct incl. fittings, joints, sealant", f"{g}, low pressure (500 Pa) class, Seal class A", "m²", a_, fl)
    for (fl, d), L in sorted(flex.items()):
        row(sc, f"Insulated flexible duct Ø{U.size(d)}", R["flex_duct"], "m", L, fl)

    # D. air terminals
    sd = sec("D. AIR TERMINALS AND LOUVERS", "air_terminals")
    term = defaultdict(list)
    for a in design.get("air_systems", []):
        for d in a["diffusers"]:
            term[("4-way square ceiling diffuser with OBD and plenum box", f"{d['model']} neck Ø{U.size(d['neck_mm'])}, face {U.rect(d['face_mm'], d['face_mm'])}", a["floor"])].append(d["tag"])
        g = a["return"]["grille"]
        term[("Return grille with filter and plenum box", f"{U.rect(g['w_mm'], g['h_mm'])}, aluminium, eggcrate", a["floor"])].append(g["tag"])
        if a.get("fresh_air"):
            lv = a["fresh_air"]["louver"]
            term[("Fresh air louver with bird screen and sand trap", f"{U.rect(lv['w_mm'], lv['h_mm'])}, aluminium", a["floor"])].append(lv["tag"])
        if a["transfer_grilles"]:
            term[("Door transfer grille", p["design"]["ducts"]["door_transfer_grille"] + ", by door supplier / HVAC", a["floor"])] += [f"DTG-{r}" for r in a["transfer_grilles"]]
    for e in design.get("exhaust_systems", []):
        g = e["grille"]
        term[("Exhaust grille", f"{U.rect(g['w_mm'], g['h_mm'])}, aluminium, eggcrate", e["floor"])].append(g["tag"])
        if e["discharge"]["type"] == "wall_louver":
            d = e["discharge"]
            term[("Exhaust wall louver / cowl with backdraft shutter", f"{U.rect(d['w_mm'], d['h_mm'])}, aluminium", e["floor"])].append(d["tag"])
    for (desc, spec, fl), tags in sorted(term.items()):
        row(sd, desc, spec, "No.", len(tags), fl, [t for t in tags if not t.startswith("DTG-")], remarks="rooms " + ", ".join(t[4:] for t in tags if t.startswith("DTG-")) if any(t.startswith("DTG-") for t in tags) else "")

    # E. dampers
    se = sec("E. DAMPERS", "dampers")
    for (fl, size), zs in sorted(vcd.items()):
        row(se, "Volume control damper (branch), opposed blade", size, "No.", len(zs), fl, remarks="zones " + ", ".join(sorted(set(zs))))
    for a in design.get("air_systems", []):
        if a.get("fresh_air"):
            d = a["fresh_air"]["segments"][0]["d_mm"]
            row(se, "Volume control damper on fresh air duct", f"Ø{U.size(d)}", "No.", 1, a["zone"])
            row(se, "Motorised fresh air damper, 24 V on/off, interlocked with indoor unit", f"Ø{U.size(d)}", "No.", 1, a["zone"])
    nb = sum(1 for e in design.get("exhaust_systems", []))
    row(se, "Backdraft (non-return) damper at exhaust discharge", "Gravity type", "No.", nb, "All")
    se["rows"].append({"desc": "Fire dampers at fire-rated walls / floors", "spec": "UL 555, 1.5 h, fusible link", "unit": "No.", "qty": 0, "zone": "TBC",
                       "remarks": R["fire_damper_note"], "tags": []})

    # F. fans
    sf = sec("F. FANS", "fans")
    fans = defaultdict(list)
    for e in design.get("exhaust_systems", []):
        f = e["fan"]
        fans[(f"Exhaust fan, {f['type']}", f"{f['model']}: {U.flow(f['airflow_ls'])} @ {f['esp_pa']} Pa, {f['power_w']} W", e["floor"])].append(f["tag"])
    for (desc, spec, fl), tags in sorted(fans.items()):
        row(sf, desc, spec, "No.", len(tags), fl, tags)

    # G. pipes and fittings
    sg_ = sec("G. REFRIGERANT / CHILLED WATER PIPING AND FITTINGS", "pipes")
    cu = defaultdict(float)
    elbows = defaultdict(int)
    for r in design.get("pipes", {}).get("refrigerant", []):
        for od, inch in ((r["liquid_mm"], r["liquid_in"]), (r["gas_mm"], r["gas_in"])):
            cu[(od, inch)] += r["total_m"]
            elbows[(od, inch)] += R["refrigerant_fittings_per_run"]["elbows_fixed"] + r["bends"] * R["refrigerant_fittings_per_run"]["elbows_per_bend"]
    for (od, inch), L in sorted(cu.items()):
        row(sg_, "Copper refrigerant pipe, ACR grade, seamless, deoxidised", f"OD {od} mm ({inch}\")", "m", L, "All", remarks="riser + floor + roof runs")
    for (od, inch), n in sorted(elbows.items()):
        row(sg_, "Copper refrigerant fittings (long-radius elbows / couplings), brazed with nitrogen purge", f"OD {od} mm", "No.", n, "All", remarks=None or "")
    kits = design.get("pipes", {}).get("branch_kits", [])
    row(sg_, "VRF refnet branch joint kit", "Per brand selection", "No.", len(kits), "All", [k["unit"] for k in kits])
    for c in design.get("pipes", {}).get("chw", []):
        row(sg_, "Black steel pipe Sch.40 (CHW supply + return)", f"DN{c['dn']} ({c['in']}\")", "m", c["length_m"], c["floor"], [c["tag"]])

    # H. insulation
    sh = sec("H. THERMAL INSULATION", "insulation")
    for (fl, spec), a_ in sorted(ins_area.items()):
        row(sh, "Duct insulation (supply / return / fresh air)", spec, "m²", a_, fl)
    ins = PS["refrigerant"]["insulation_mm"]
    for (od, inch), L in sorted(cu.items()):
        t = ins["gas"] if any(od == r["gas_mm"] for r in design.get("pipes", {}).get("refrigerant", [])) else ins["liquid"]
        row(sh, "Closed-cell elastomeric pipe insulation, UV-protected outdoors", f"for OD {od} mm, {t} mm thick", "m", L, "All")
    for c in design.get("pipes", {}).get("chw", []):
        t = next(th for lim, th in PS["chilled_water"]["insulation_mm"] if c["dn"] <= lim)
        row(sh, "CHW pipe insulation with vapour barrier and cladding outdoors", f"DN{c['dn']}, {t} mm", "m", c["length_m"], c["floor"])

    # I. condensate
    si = sec("I. CONDENSATE DRAIN", "condensate")
    cd = defaultdict(list)
    for c in design.get("pipes", {}).get("condensate", []):
        cd[(c["dn_mm"], c["floor"])].append(c)
    for (dn, fl), cs in sorted(cd.items()):
        row(si, "uPVC condensate pipe incl. fittings, slope 1 %", f"Ø{U.size(dn)}, {PS['condensate']['material']}", "m", sum(x["length_m"] for x in cs), fl, [x["tag"] for x in cs])
        row(si, "Condensate trap / P-trap with cleanout at unit", f"Ø{U.size(dn)}", "No.", len(cs), fl)

    # J. supports
    sj = sec("J. SUPPORTS AND HANGERS", "supports")
    sp = R["duct_supports_spacing_m"].get(std, 2.4)
    for fl, L in sorted(rigid_len.items()):
        row(sj, "Duct hangers: threaded rod + angle/strap, galvanised", f"@ {sp} m max (pairs)", "No.", math.ceil(L / sp), fl)
    tot_cu = sum(r["total_m"] for r in design.get("pipes", {}).get("refrigerant", []))
    row(sj, "Refrigerant pipe supports / clamps with insulation saddles", f"@ {PS['refrigerant']['support_spacing_m']['le_15.88']} m", "No.", math.ceil(tot_cu / PS["refrigerant"]["support_spacing_m"]["le_15.88"]), "All")
    tot_cd = sum(c["length_m"] for c in design.get("pipes", {}).get("condensate", []))
    row(sj, "Condensate pipe clips", f"@ {PS['condensate']['support_spacing_m']} m", "No.", math.ceil(tot_cd / PS["condensate"]["support_spacing_m"]), "All")
    row(sj, "Outdoor unit steel support frame / concrete plinth with anti-vibration pads", "Galvanised, by structural check", "No.", len(design["outdoor_units"]), "Roof")

    # K. controls
    sk = sec("K. CONTROLS", "controls")
    for z in design["zones"]:
        t = z["thermostat"]
        row(sk, "Wired programmable thermostat / controller" if z["kind"] == "ducted" else "Wired room controller", t["type"], "No.", 1, z["id"], [t["tag"]])
    n_md = sum(1 for a in design.get("air_systems", []) if a.get("fresh_air"))
    row(sk, "Control wiring, interlocks (FA damper / FA fan to unit), communication cable", "Shielded 2-core, in conduit", "Lot", 1 if design["zones"] else 0, "All", remarks=f"{n_md} motorised damper interlocks")

    # L. T&C
    sl = sec("L. TESTING AND COMMISSIONING", "testing")
    for d, u in R["testing_items"]:
        row(sl, d, "", u, 1, "All")

    tags = sorted({t for s_ in sections for r_ in s_["rows"] for t in r_["tags"]})
    C.save(work, "boq.json", {"sections": sections, "tags": tags})
    write_xlsx(work, p, sections, W)
    n = sum(len(s_["rows"]) for s_ in sections)
    print(f"boq: {n} line items in {len(sections)} sections -> out/HVAC_BOQ.xlsx")


def write_xlsx(work, p, sections, W):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    note = p["design"]["drawing"]["preliminary_note"]
    thin = Side(style="thin", color="999999")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = f"BILL OF QUANTITIES - HVAC (QUANTITIES ONLY, NO PRICES) - {p['project']['name']}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = note
    ws["A2"].font = Font(bold=True, color="C00000")
    ws["A3"] = f"Location: {p['project']['location']['city']}, {p['project']['location']['country']}   Codes: {p['codes']['label']}   Units: {p['units']}"
    hdr = ["Item No.", "Description", "Specification", "Unit", "Quantity", "Waste %", "Qty incl. allowance", "Floor/Zone", "Remarks"]
    r = 5
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = bd
    ws.freeze_panes = "A6"
    r += 1
    for s in sections:
        if not s["rows"]:
            continue
        ws.cell(row=r, column=1, value=s["title"]).font = Font(bold=True)
        for c in range(1, 10):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="DDEBF7")
        r += 1
        letter = s["title"].split(".")[0]
        for i, it in enumerate(s["rows"], 1):
            vals = [f"{letter}.{i}", it["desc"], it["spec"], it["unit"], it["qty"], W.get(s["key"], 0), None, it["zone"], it["remarks"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = bd
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=7, value=f"=ROUND(E{r}*(1+F{r}/100),2)" if it["unit"] not in ("No.", "Lot") else f"=ROUNDUP(E{r}*(1+F{r}/100),0)").border = bd
            ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor="FFF2CC")
            r += 1
        r += 1
    for c, w in zip(range(1, 10), [8, 46, 46, 7, 10, 9, 12, 12, 50]):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws2 = wb.create_sheet("Notes")
    lines = [note, "", "Quantities are measured from the preliminary design model (design.json) and include the waste/allowance % shown (editable, yellow cells).",
             "Ductwork measured as external surface area of straight duct by gauge; fittings covered by the allowance.",
             "Equipment and terminals are generic sizes - replace with the approved brand's model numbers.",
             "Fire dampers, builder's work openings, electrical power wiring and BMS integration are excluded / TBC.", "",
             "See assumptions.md and the design calculation report for the design basis."]
    for i, t in enumerate(lines, 1):
        ws2.cell(row=i, column=1, value=t)
    ws2.column_dimensions["A"].width = 140
    os.makedirs(os.path.join(work, "out"), exist_ok=True)
    wb.save(os.path.join(work, "out", "HVAC_BOQ.xlsx"))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-w", "--work", required=True)
    build(ap.parse_args().work)
