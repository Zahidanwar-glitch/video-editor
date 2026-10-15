"""Unit + end-to-end tests.  Run:  python -m unittest discover -s tests -v   (from the skill folder)"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(SKILL, "scripts"))

import common as C  # noqa: E402
import duct_sizing as DSZ  # noqa: E402
import pipe_sizing as PS  # noqa: E402

EX = os.path.join(SKILL, "examples", "riyadh_villa")


class Calcs(unittest.TestCase):
    def setUp(self):
        DSZ.DS = C.ref("duct_sizing.json")

    def test_duct_friction_matches_ashrae_chart(self):
        fr, v = DSZ.fric(500, 400)          # 500 L/s in 400 mm round: chart ~0.45 Pa/m, 4.0 m/s
        self.assertAlmostEqual(v, 3.98, places=1)
        self.assertTrue(0.35 < fr < 0.6, fr)

    def test_equivalent_diameter(self):
        self.assertAlmostEqual(DSZ.de_rect(400, 300), 378, delta=3)   # ASHRAE Table: 400x300 -> 378 mm

    def test_rect_sizing_respects_depth_and_velocity(self):
        w, h, de, fr, v, ok = DSZ.size_rect(750, 0.8, 5.0, 300, 50, 150, 4.0)
        self.assertLessEqual(h, 300)
        self.assertLessEqual(v, 5.0)
        self.assertLessEqual(fr, 0.85)

    def test_psychrometrics(self):
        self.assertAlmostEqual(C.w_from_rh(24, 50, 101.325) * 1000, 9.3, delta=0.2)
        self.assertAlmostEqual(C.p_atm_kpa(612), 94.2, delta=0.2)

    def test_refrigerant_and_chw_tables(self):
        row = PS.refrig_size(7.0, C.ref("pipe_sizing.json")["refrigerant"]["by_capacity_kw"])
        self.assertEqual((row["liquid_mm"], row["gas_mm"]), (9.52, 15.88))
        s, v, dp = PS.chw_size(1.0, C.ref("pipe_sizing.json")["chilled_water"])
        self.assertLessEqual(dp, 400)
        self.assertLessEqual(v, 1.2)

    def test_code_lookup(self):
        self.assertEqual(C.code_set_for("Saudi Arabia"), "KSA")
        self.assertEqual(C.ref("codes_by_country.json")["code_sets"]["USA"]["units"], "IP")
        self.assertEqual(C.code_set_for("Narnia"), "DEFAULT")
        self.assertIsNotNone(C.climate_for("Riyadh", "Saudi Arabia"))

    def test_scale_detection_on_sample(self):
        import pymupdf
        import scale_detect
        r = scale_detect.detect(pymupdf.open(os.path.join(EX, "sample_villa.pdf"))[0])
        self.assertEqual(r["denominator"], 100)
        self.assertEqual(r["confidence"], "high")


class EndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.work = tempfile.mkdtemp(prefix="hvac_e2e_")
        shutil.copy(os.path.join(EX, "project.json"), cls.work)
        r = subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "pipeline.py"), "all", "-w", cls.work,
                            "--pdf", os.path.join(EX, "sample_villa.pdf")], capture_output=True, text=True)
        cls.log = r.stdout + r.stderr
        assert r.returncode == 0, cls.log

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.work, ignore_errors=True)

    def j(self, name):
        with open(os.path.join(self.work, name)) as f:
            return json.load(f)

    def test_rooms_read(self):
        rooms = self.j("rooms.json")
        names = {r["name"] for f in rooms["floors"] for r in f["rooms"]}
        for n in ("MAJLIS", "LIVING", "KITCHEN", "MASTER BEDROOM", "BEDROOM 4", "SHAFT"):
            self.assertIn(n, names)
        gf = next(f for f in rooms["floors"] if f["id"] == "GF")
        majlis = next(r for r in gf["rooms"] if r["name"] == "MAJLIS")
        self.assertAlmostEqual(majlis["area_m2"], 27.31, delta=27.31 * 0.05)
        self.assertEqual({e["orientation"] for e in majlis["exterior_walls"]}, {"N", "W"})
        self.assertEqual({w["tag"] for w in majlis["windows"]}, {"W1", "W2"})
        self.assertAlmostEqual(gf["north_deg"], 20, delta=1.5)

    def test_every_room_and_check(self):
        checks = self.j("checks.json")["checks"]
        self.assertFalse([c for c in checks if c["status"] == "FAIL"], checks)

    def test_deliverables_and_note(self):
        out = os.path.join(self.work, "out")
        for f in ("HVAC_plans.pdf", "HVAC_GF.dxf", "HVAC_FF.dxf", "HVAC_BOQ.xlsx", "HVAC_Design_Report.docx", "assumptions.md"):
            self.assertTrue(os.path.exists(os.path.join(out, f)), f)
        note = C.ref("design_defaults.json")["drawing"]["preliminary_note"]
        import pymupdf
        from docx import Document
        from openpyxl import load_workbook
        for pg in pymupdf.open(os.path.join(out, "HVAC_plans.pdf")):
            self.assertIn("PRELIMINARY DESIGN", pg.get_text())
        self.assertIn(note, Document(os.path.join(out, "HVAC_Design_Report.docx")).sections[0].footer.paragraphs[0].text)
        self.assertEqual(load_workbook(os.path.join(out, "HVAC_BOQ.xlsx"))["BOQ"]["A2"].value, note)

    def test_dxf_layers(self):
        import ezdxf
        doc = ezdxf.readfile(os.path.join(self.work, "out", "HVAC_GF.dxf"))
        used = {e.dxf.layer for e in doc.modelspace()}
        for lay in ("M-DUCT-SUPP", "M-DUCT-RETN", "M-DUCT-FRSH", "M-DUCT-EXHS", "M-PIPE-REFG", "M-PIPE-COND", "M-EQPM", "M-CTRL", "M-ANNO"):
            self.assertIn(lay, used)

    def test_mixed_system(self):
        d = self.j("design.json")
        systems = {z["rooms"][0]: z["system"] for z in d["zones"]}
        self.assertEqual(systems.get("GF-06"), "wall_split")   # kitchen
        self.assertTrue(any(z["system"] == "ducted_split" for z in d["zones"]))
        for z in d["zones"]:
            self.assertGreaterEqual(z["unit"]["derated_kw"], z["unit"]["req_total_kw"])


if __name__ == "__main__":
    unittest.main()
