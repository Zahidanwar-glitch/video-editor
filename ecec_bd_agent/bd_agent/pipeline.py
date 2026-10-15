"""Daily run: ingest -> qualify -> score -> rank -> capture plans -> reports -> database."""
from __future__ import annotations

import json
from datetime import date, timedelta
from pathlib import Path

from . import capture, learning, reports, scoring
from .store import (DATA, OUTPUT, TERMINAL, STATUSES, audit, config_hash, find_duplicate, load_config,
                    load_db, load_json, load_kb, now_ast, opp_id, save_db, save_json)

CORE_FIELDS = ["project", "client", "country", "location", "sector", "opportunity_type", "services_needed",
               "estimated_project_value_usd", "expected_consulting_value_usd", "current_stage", "bid_due_date",
               "project_duration_months", "procurement_route", "main_contractor", "consultant", "decision_makers",
               "source", "source_date", "description", "missing", "analyst"]


# ------------------------------------------------------------------ status sync

def sync_status_from_excel(db: dict) -> list[str]:
    """Pick up Status edits the BD team made directly in opportunities.xlsx."""
    path = OUTPUT / "opportunities.xlsx"
    if not path.exists():
        return []
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    if "Opportunity Register" not in wb.sheetnames:
        return []
    ws = wb["Opportunity Register"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    hdr = [str(h) for h in rows[0]]
    i_id, i_st = hdr.index("ID"), hdr.index("Status")
    changed = []
    for r in rows[1:]:
        oid, st = r[i_id], (str(r[i_st]).strip().upper() if r[i_st] else None)
        opp = db["opportunities"].get(oid)
        # only honour cells a person edited since the agent last wrote the workbook
        if opp and st and st in STATUSES and st != opp["status"] and st != opp.get("excel_status"):
            changed.append(set_status(db, oid, st, by="excel"))
    return changed


def set_status(db: dict, oid: str, status: str, by: str = "cli", note: str = "") -> str:
    status = status.upper()
    if status not in STATUSES:
        raise ValueError(f"Unknown status {status}; use one of {STATUSES}")
    opp = db["opportunities"][oid]
    old = opp["status"]
    opp["status"] = status
    opp.setdefault("status_history", []).append(
        {"ts": now_ast().isoformat(timespec="seconds"), "from": old, "to": status, "by": by, "note": note})
    audit("status_change", id=oid, project=opp["project"], frm=old, to=status, by=by, note=note)
    if status in ("WON", "LOST"):
        learning.record_outcome(opp, status, note)
    return f"{oid}: {old} -> {status}"


# ------------------------------------------------------------------ ingest

def ingest(db: dict, inbox_files: list[Path], cfg: dict, today: date) -> dict:
    stats = {"received": 0, "new": 0, "updated": 0, "out_of_scope": 0, "intel": 0}
    mi = load_json(DATA / "market_intelligence.json", [])
    seen_mi = {(m["headline"].lower()[:80]) for m in mi}
    for f in inbox_files:
        payload = json.loads(Path(f).read_text(encoding="utf-8"))
        items = payload.get("opportunities", payload if isinstance(payload, list) else [])
        for it in items:
            if "market_intelligence" in it:
                payload.setdefault("market_intelligence", []).extend(it["market_intelligence"])
                continue
            stats["received"] += 1
            if it.get("country") not in cfg["gcc_countries"]:
                stats["out_of_scope"] += 1
                audit("ignored", reason="outside target GCC countries", project=it.get("project"), country=it.get("country"))
                continue
            dup = find_duplicate(it, {**db["opportunities"], **db["rejected"]})
            if dup:
                rec = db["opportunities"].get(dup) or db["rejected"][dup]
                for k in CORE_FIELDS:
                    if it.get(k) not in (None, "", []) and it.get(k) != rec.get(k):
                        rec.setdefault("change_log", []).append({"date": today.isoformat(), "field": k, "old": rec.get(k), "new": it[k]})
                        rec[k] = it[k]
                rec["last_seen"] = today.isoformat()
                srcs = rec.setdefault("sources", [])
                if it.get("source") and it["source"] not in srcs:
                    srcs.append(it["source"])
                stats["updated"] += 1
                audit("duplicate_merged", id=dup, project=it.get("project"), source=it.get("source"))
                continue
            oid = opp_id(it)
            rec = {k: it.get(k) for k in CORE_FIELDS}
            rec.update(id=oid, first_seen=today.isoformat(), last_seen=today.isoformat(), status="NEW",
                       sources=[it.get("source")] if it.get("source") else [], discovered_by=payload.get("researcher"))
            db["opportunities"][oid] = rec
            stats["new"] += 1
            audit("discovered", id=oid, project=rec["project"], client=rec["client"], source=rec["source"])
        for m in payload.get("market_intelligence", []):
            key = m["headline"].lower()[:80]
            if key not in seen_mi:
                m["logged"] = today.isoformat()
                mi.append(m)
                seen_mi.add(key)
                stats["intel"] += 1
    save_json(DATA / "market_intelligence.json", mi)
    return stats


# ------------------------------------------------------------------ scoring

def score_all(db: dict, cfg: dict, kb_raw: dict, today: date) -> dict:
    kb = scoring.KBIndex(kb_raw)
    cal = learning.calibration(cfg)
    chash = config_hash()
    active = {k: v for k, v in db["opportunities"].items() if v["status"] not in TERMINAL}
    go_count = sum(1 for v in active.values() if v.get("decision", {}).get("recommendation") == "GO")
    rejected = 0
    for oid, opp in list(active.items()):
        fit = scoring.fit_model(opp, kb, cfg, go_count)
        if fit["technical_alignment_pct"] < cfg["min_technical_alignment_pct"] and opp["status"] in ("NEW", "QUALIFIED"):
            opp.update(fit=fit, status="REJECTED", rejected_reason=f"Technical alignment {fit['technical_alignment_pct']}% < {cfg['min_technical_alignment_pct']}%")
            db["rejected"][oid] = db["opportunities"].pop(oid)
            audit("rejected", id=oid, project=opp["project"], reason=opp["rejected_reason"], model=cfg["model_version"], config=chash)
            rejected += 1
            continue
        risk = scoring.risk_model(opp, fit, kb, cfg, today)
        win = scoring.win_model(opp, fit, risk, kb, cfg, cal["shift"])
        dec = scoring.decision(opp, fit, risk, win, cfg)
        prev = opp.get("decision", {}).get("recommendation")
        opp.update(fit=fit, risk=risk, win=win, decision=dec,
                   priority_score=scoring.priority_score(fit, risk, win, dec),
                   missing_info=scoring.missing_info(opp, fit), scored_on=today.isoformat())
        opp["actions"] = capture.actions(opp, today)
        if opp["status"] == "NEW" and opp["first_seen"] != today.isoformat():
            opp["status"] = "QUALIFIED"
        audit("scored", id=oid, project=opp["project"], model=cfg["model_version"], config=chash,
              fit=fit["score"], risk=risk["score"], win=win["probability"], recommendation=dec["recommendation"],
              previous_recommendation=prev,
              fit_criteria={k: [v["score"], v["method"], v["why"]] for k, v in fit["criteria"].items()},
              risk_criteria={k: [v["score"], v["method"]] for k, v in risk["categories"].items()},
              win_factors={k: v["value"] for k, v in win["factors"].items()}, reasons=dec["reasons"])
    ranked = sorted((o for o in db["opportunities"].values() if o["status"] not in TERMINAL),
                    key=lambda o: -o["priority_score"])
    for i, o in enumerate(ranked, 1):
        o["priority_rank"] = i
    return {"calibration": cal, "rejected": rejected, "scored": len(ranked)}


def auto_close(db: dict, today: date, grace_days: int = 14) -> list[str]:
    closed = []
    for oid, o in db["opportunities"].items():
        if o["status"] in ("NEW", "QUALIFIED") and o.get("bid_due_date"):
            try:
                due = date.fromisoformat(o["bid_due_date"][:10])
            except ValueError:
                continue
            if due < today - timedelta(days=grace_days):
                closed.append(set_status(db, oid, "CLOSED", by="agent", note=f"Due date {due} passed > {grace_days} days, no pursuit"))
    return closed


def run(inbox_files: list[Path], today: date | None = None, weekly: bool | None = None) -> dict:
    today = today or now_ast().date()
    cfg, kb, db = load_config(), load_kb(), load_db()
    audit("run_start", date=today.isoformat(), inbox=[str(p) for p in inbox_files], model=cfg["model_version"], config=config_hash())
    excel_changes = sync_status_from_excel(db)
    stats = ingest(db, inbox_files, cfg, today)
    closed = auto_close(db, today)
    sstats = score_all(db, cfg, kb, today)
    plans = []
    for o in db["opportunities"].values():
        if o["status"] in ("APPROVED", "IN PURSUIT"):
            plans.append(capture.write_capture_plan(o, cfg, kb, today))
            o["capture_plan"] = plans[-1]
    for o in list(db["opportunities"].values()) + list(db["rejected"].values()):
        o["excel_status"] = o["status"]
    save_db(db)
    if weekly is None:
        weekly = today.weekday() == 6  # Sunday = start of the Saudi working week
    out = reports.build_all(db, cfg, kb, today, {**stats, **sstats, "excel_changes": excel_changes,
                                                  "auto_closed": closed, "capture_plans": plans}, weekly)
    audit("run_end", date=today.isoformat(), **{k: v for k, v in stats.items()}, scored=sstats["scored"],
          rejected=sstats["rejected"], outputs=out)
    return {"stats": stats, "scoring": sstats, "excel_changes": excel_changes, "auto_closed": closed,
            "capture_plans": plans, "outputs": out}
