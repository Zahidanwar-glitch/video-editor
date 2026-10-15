"""Outputs: opportunities.xlsx (register + risk + intel + approved + audit), daily Markdown and PDF report,
weekly trend analysis."""
from __future__ import annotations

from collections import Counter
from datetime import date, timedelta

from .store import DATA, OUTPUT, TERMINAL, STATUSES, load_json

REC_ORDER = {"GO": 0, "GO WITH CAUTION": 1, "NO GO": 2}


def _m(v):
    return f"{v / 1e6:,.1f}" if v else "n/a"


def _active(db):
    return [o for o in db["opportunities"].values() if o["status"] not in TERMINAL and "fit" in o]


def _days(o, today):
    try:
        return (date.fromisoformat(o["bid_due_date"][:10]) - today).days
    except (TypeError, ValueError):
        return None


# ------------------------------------------------------------------ Excel

def write_excel(db, cfg, today, mi, path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    hdr_fill = PatternFill("solid", fgColor="1F3A5F")
    hdr_font = Font(bold=True, color="FFFFFF")
    wb = Workbook()

    def sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(40, len(h) + 4)))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    allopps = sorted(db["opportunities"].values(), key=lambda o: (o["status"] in TERMINAL, o.get("priority_rank", 999)))
    # Sheet 1
    ws = wb.active
    ws.title = "Opportunity Register"
    H = ["ID", "Priority Ranking", "Date", "Project", "Client", "Country", "Sector", "Opportunity Type",
         "Project Value (USD M)", "Consultancy Value (USD M)", "Fit Score", "Risk Score", "Risk Class",
         "Win Probability", "Recommendation", "Status", "Bid Due Date", "Expected Value (USD M)",
         "Return on Effort (x)", "Missing / Uncertain Info", "Source"]
    rows = []
    for o in allopps:
        if "fit" not in o or "risk" not in o:
            continue
        rows.append([o["id"], o.get("priority_rank"), o["first_seen"], o["project"], o["client"], o["country"], o["sector"],
                     o.get("opportunity_type"), (o.get("estimated_project_value_usd") or 0) / 1e6 or None,
                     (o["fit"]["consulting_value_usd"] or 0) / 1e6 or None, o["fit"]["score"], o["risk"]["score"],
                     o["risk"]["classification"], o["win"]["probability"], o["decision"]["recommendation"], o["status"],
                     o.get("bid_due_date"), o["decision"]["expected_value_usd"] / 1e6, o["decision"]["return_on_effort"],
                     "; ".join(o.get("missing_info", []))[:400], o.get("source")])
    sheet(ws, H, rows, {"Project": 48, "Client": 30, "Missing / Uncertain Info": 45, "Source": 40, "Recommendation": 18})
    n = len(rows) + 1
    for col in ("I", "J", "R"):
        for c in ws[col][1:]:
            c.number_format = "#,##0.0"
    for c in ws["N"][1:]:
        c.number_format = "0%"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    dv.error = "Choose a valid status"
    ws.add_data_validation(dv)
    dv.add(f"P2:P{max(n, 2)}")
    green, amber, red = (PatternFill("solid", fgColor=c) for c in ("C6EFCE", "FFEB9C", "FFC7CE"))
    ws.conditional_formatting.add(f"O2:O{n}", CellIsRule(operator="equal", formula=['"GO"'], fill=green))
    ws.conditional_formatting.add(f"O2:O{n}", CellIsRule(operator="equal", formula=['"GO WITH CAUTION"'], fill=amber))
    ws.conditional_formatting.add(f"O2:O{n}", CellIsRule(operator="equal", formula=['"NO GO"'], fill=red))
    ws.conditional_formatting.add(f"L2:L{n}", CellIsRule(operator="greaterThan", formula=["60"], fill=red))

    # Sheet 2
    ws = wb.create_sheet("Detailed Risk Register")
    cats = cfg["risk_categories"]
    H = ["ID", "Project", "Country"] + [c.replace("_", " ").title() for c in cats] + ["Overall Risk", "Classification", "Top Risks", "Justification (top risks)"]
    rows = []
    for o in allopps:
        if "risk" not in o:
            continue
        r = o["risk"]
        rows.append([o["id"], o["project"], o["country"]] + [r["categories"][c]["score"] for c in cats] +
                    [r["score"], r["classification"], ", ".join(r["top_risks"]),
                     " | ".join(f"{c}: {r['categories'][c]['why']}" for c in r["top_risks"])])
    sheet(ws, H, rows, {"Project": 45, "Justification (top risks)": 80})
    for col in range(4, 4 + len(cats)):
        L = get_column_letter(col)
        ws.conditional_formatting.add(f"{L}2:{L}{len(rows) + 1}", CellIsRule(operator="greaterThanOrEqual", formula=["4"], fill=red))

    # Sheet 3
    ws = wb.create_sheet("Market Intelligence")
    sheet(ws, ["Date", "Logged", "Country", "Category", "Headline", "Implication for ECEC", "Source"],
          [[m.get("date"), m.get("logged"), m.get("country"), m.get("category"), m.get("headline"), m.get("implication"), m.get("source")]
           for m in sorted(mi, key=lambda m: m.get("date") or "", reverse=True)],
          {"Headline": 60, "Implication for ECEC": 60, "Source": 40})

    # Sheet 4
    ws = wb.create_sheet("Approved Opportunities")
    appr = [o for o in allopps if o["status"] in ("APPROVED", "IN PURSUIT", "SUBMITTED", "WON", "LOST")]
    sheet(ws, ["ID", "Project", "Client", "Country", "Status", "Approved On", "Fit", "Risk", "Win", "Bid Due", "Capture Plan", "Next Action", "Owner", "Target"],
          [[o["id"], o["project"], o["client"], o["country"], o["status"],
            next((h["ts"][:10] for h in reversed(o.get("status_history", [])) if h["to"] == "APPROVED"), None),
            o.get("fit", {}).get("score"), o.get("risk", {}).get("score"), o.get("win", {}).get("probability"), o.get("bid_due_date"),
            o.get("capture_plan"), (o.get("actions") or [{}])[0].get("action"), (o.get("actions") or [{}])[0].get("responsible"),
            (o.get("actions") or [{}])[0].get("target_date")] for o in appr],
          {"Project": 45, "Next Action": 60, "Capture Plan": 35})

    # Extra: scoring audit + rejected
    ws = wb.create_sheet("Scoring Audit")
    rows = []
    for o in allopps:
        if "fit" not in o or "risk" not in o:
            continue
        for k, v in o["fit"]["criteria"].items():
            rows.append([o["id"], o["project"], "Fit", k, v["score"], cfg["fit_weights"][k], v["method"], v["why"]])
        for k, v in o["risk"]["categories"].items():
            rows.append([o["id"], o["project"], "Risk", k, v["score"], None, v["method"], v["why"]])
        for k, v in o["win"]["factors"].items():
            rows.append([o["id"], o["project"], "Win factor", k, v["value"], None, "rule/analyst", v["why"]])
        rows.append([o["id"], o["project"], "Decision", o["decision"]["recommendation"], None, None, "engine", "; ".join(o["decision"]["reasons"])])
    sheet(ws, ["ID", "Project", "Model", "Criterion", "Score", "Weight", "Method", "Justification"], rows,
          {"Project": 40, "Justification": 90, "Criterion": 28})
    ws = wb.create_sheet("Rejected")
    sheet(ws, ["ID", "Date", "Project", "Client", "Country", "Reason", "Source"],
          [[o["id"], o["first_seen"], o["project"], o["client"], o["country"], o.get("rejected_reason"), o.get("source")]
           for o in db["rejected"].values()], {"Project": 45, "Reason": 50})
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ report content

def report_data(db, cfg, today, mi, run_stats):
    act = _active(db)
    ranked = sorted(act, key=lambda o: o.get("priority_rank", 999))
    new = [o for o in ranked if o["first_seen"] == today.isoformat()]
    high_risk = sorted([o for o in act if o["risk"]["score"] > 60], key=lambda o: -o["risk"]["score"])
    upcoming = sorted([o for o in act if (d := _days(o, today)) is not None and 0 <= d <= 30], key=lambda o: o["bid_due_date"])
    approved = [o for o in db["opportunities"].values() if o["status"] in ("APPROVED", "IN PURSUIT")]
    recent_mi = sorted([m for m in mi if (m.get("date") or "") >= (today - timedelta(days=7)).isoformat() or m.get("logged") == today.isoformat()],
                       key=lambda m: m.get("date") or "", reverse=True)
    acts = []
    for o in ranked[:12]:
        if o["decision"]["recommendation"] == "NO GO":
            continue
        for a in o.get("actions", [])[:2]:
            acts.append({**a, "project": o["project"], "id": o["id"], "rank": o["priority_rank"]})
    acts.sort(key=lambda a: (a["target_date"], a["rank"]))
    pipe_fee = sum(o["fit"]["consulting_value_usd"] or 0 for o in act if o["decision"]["recommendation"] != "NO GO")
    pipe_ev = sum(o["decision"]["expected_value_usd"] for o in act if o["decision"]["recommendation"] != "NO GO")
    recs = Counter(o["decision"]["recommendation"] for o in act)
    return dict(today=today, ranked=ranked, new=new, high_risk=high_risk, upcoming=upcoming, approved=approved,
                recent_mi=recent_mi, actions=acts[:20], pipe_fee=pipe_fee, pipe_ev=pipe_ev, recs=recs, stats=run_stats,
                unknown_fee=sum(1 for o in act if o["fit"]["consulting_value_usd"] is None))


def weekly_trends(db, mi, today):
    wk, prev = today - timedelta(days=7), today - timedelta(days=14)
    cur_mi = [m for m in mi if (m.get("date") or "") >= wk.isoformat()]
    old_mi = [m for m in mi if prev.isoformat() <= (m.get("date") or "") < wk.isoformat()]
    allo = list(db["opportunities"].values()) + list(db["rejected"].values())
    cur_o = [o for o in allo if o["first_seen"] >= wk.isoformat()]
    act = _active(db)
    return {
        "mi_by_category": Counter(m.get("category") for m in cur_mi),
        "mi_by_country": Counter(m.get("country") for m in cur_mi),
        "mi_delta": len(cur_mi) - len(old_mi),
        "new_opps_by_sector": Counter(o.get("sector") for o in cur_o),
        "new_opps_by_country": Counter(o.get("country") for o in cur_o),
        "pipeline_by_sector_ev": {s: sum(o["decision"]["expected_value_usd"] for o in act if o["sector"] == s)
                                  for s in {o["sector"] for o in act}},
        "go_rate": round(sum(o["decision"]["recommendation"] != "NO GO" for o in act) / len(act), 2) if act else None,
    }


def _row_md(o, today):
    d = _days(o, today)
    due = f"{o['bid_due_date']} ({d}d)" if d is not None else "TBC"
    return (f"| {o.get('priority_rank')} | {o['project'][:70]} | {o['client'][:35]} | {o['country']} | "
            f"{o['fit']['score']:.0f} | {o['risk']['score']:.0f} | {o['win']['probability']:.0%} | "
            f"{_m(o['fit']['consulting_value_usd'])} | **{o['decision']['recommendation']}** | {due} |")


def markdown(rd, cfg, cal, trends):
    t = rd["today"]
    s = rd["stats"]
    L = [f"# ECEC BD Intelligence - Daily Executive Report", f"**Date:** {t.isoformat()} (AST) | **Model:** v{cfg['model_version']} | "
         f"**Pipeline (GO + caution):** USD {_m(rd['pipe_fee'])}M fees, USD {_m(rd['pipe_ev'])}M risk-weighted", "",
         "## Executive Summary",
         f"- **{s.get('new', 0)} new** opportunities logged, {s.get('updated', 0)} updated, {s.get('rejected', 0)} rejected on technical alignment, "
         f"{s.get('out_of_scope', 0)} outside the target GCC countries.",
         f"- Active register: **{rd['recs'].get('GO', 0)} GO**, **{rd['recs'].get('GO WITH CAUTION', 0)} GO WITH CAUTION**, "
         f"**{rd['recs'].get('NO GO', 0)} NO GO**.",
         f"- **{len(rd['upcoming'])}** submissions due in the next 30 days; **{len(rd['high_risk'])}** high-risk items.",
         f"- Data quality: fee unknown on {rd['unknown_fee']} active items, so their economics are provisional (flagged in each row).",
         ]
    if s.get("excel_changes"):
        L.append(f"- Status changes picked up from Excel: {'; '.join(s['excel_changes'])}")
    if rd["ranked"]:
        top = rd["ranked"][0]
        L.append(f"- **Best return on effort:** {top['project']} ({top['client']}): {top['decision']['recommendation']}, "
                 f"win {top['win']['probability']:.0%}, ROE {top['decision']['return_on_effort']}x.")
    TH = "| # | Project | Client | Ctry | Fit | Risk | Win | Fee $M | Rec. | Due |\n|---|---|---|---|---|---|---|---|---|---|"
    L += ["", "## 1. New Opportunities", TH] + [_row_md(o, t) for o in rd["new"]] + (["_None today._"] if not rd["new"] else [])
    L += ["", "## 2. Top 20 Priority Opportunities (ranked by expected return on effort)", TH] + [_row_md(o, t) for o in rd["ranked"][:20]]
    L += ["", "### Why the top 5 rank where they do"]
    for o in rd["ranked"][:5]:
        L.append(f"- **{o['project']}**: {'; '.join(o['decision']['reasons'])}. _Win rationale:_ {o['win']['rationale']}")
    L += ["", "## 3. High-Risk Opportunities (risk > 60)"]
    L += [f"- **{o['project']}**: risk {o['risk']['score']} ({o['risk']['classification']}). Top risks: " +
          "; ".join(f"{c} {o['risk']['categories'][c]['score']}/5: {o['risk']['categories'][c]['why']}" for c in o["risk"]["top_risks"])
          for o in rd["high_risk"]] or ["_No active opportunity scores above 60. Highest-risk categories by item are in the Excel risk register._"]
    L += ["", "## 4. Upcoming Deadlines (next 30 days)", "| Due | Days | Project | Client | Rec. | Status |", "|---|---|---|---|---|---|"]
    L += [f"| {o['bid_due_date']} | {_days(o, t)} | {o['project'][:70]} | {o['client'][:35]} | {o['decision']['recommendation']} | {o['status']} |" for o in rd["upcoming"]]
    L += ["", "## 5. Approved Opportunities"]
    L += [f"- **{o['project']}**: status {o['status']}; capture plan `{o.get('capture_plan')}`" for o in rd["approved"]] or \
         ["_None yet. Set Status = APPROVED in `output/opportunities.xlsx` (or run `python -m bd_agent set-status <ID> APPROVED`) and the next run will generate a capture plan._"]
    L += ["", "## 6. Major Market Developments (last 7 days)"]
    L += [f"- **{m.get('date')} | {m.get('country')} | {m.get('category')}:** {m.get('headline')}. _Implication:_ {m.get('implication')} ([source]({m.get('source')}))"
          for m in rd["recent_mi"][:15]]
    if trends:
        L += ["", "### Weekly Trend Analysis",
              f"- Intel items this week by category: {dict(trends['mi_by_category'])} (change vs prior week: {trends['mi_delta']:+d})",
              f"- By country: {dict(trends['mi_by_country'])}",
              f"- New opportunities by sector: {dict(trends['new_opps_by_sector'])}; by country: {dict(trends['new_opps_by_country'])}",
              f"- Risk-weighted pipeline by sector (USD M): {({k: round(v / 1e6, 1) for k, v in sorted(trends['pipeline_by_sector_ev'].items(), key=lambda kv: -kv[1])})}",
              f"- Share of active items not rated NO GO: {trends['go_rate']}"]
    L += ["", "## 7. Recommended Actions", "| Target | Opp # | Action | Impact | Effort | Owner |", "|---|---|---|---|---|---|"]
    L += [f"| {a['target_date']} | {a['rank']} | {a['action']} ({a['project'][:40]}) | {a['expected_impact']} | {a['effort']} | {a['responsible']} |" for a in rd["actions"]]
    L += ["", "---", "### Model & Learning Notes",
          f"- Win-model calibration: {cal['n_outcomes']} recorded outcomes; shift applied {cal['shift']:+.3f} logit"
          + (f"; Brier {cal['brier']}" if cal["brier"] is not None else " (calibration starts after "
             f"{cfg['learning']['min_outcomes_for_calibration']} WON/LOST outcomes)."),
          "- Scores are produced by rules applied to the ECEC knowledge base, plus analyst overrides where they are labelled. Every criterion's justification is in the Excel 'Scoring Audit' sheet and `data/audit_log.jsonl`.",
          "- Fee values marked 'estimated' are derived from capex percentages and should be validated before any bid commitment."]
    return "\n".join(L) + "\n"


# ------------------------------------------------------------------ PDF

def pdf(rd, cfg, cal, trends, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    small = ParagraphStyle("s", parent=ss["BodyText"], fontSize=7.5, leading=9.5)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9, leading=12)
    h1, h2 = ss["Title"], ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor("#1F3A5F"))
    t = rd["today"]
    P = lambda x, st=small: Paragraph(str(x).replace("&", "&amp;").replace("<", "&lt;"), st)
    rc = {"GO": colors.HexColor("#C6EFCE"), "GO WITH CAUTION": colors.HexColor("#FFEB9C"), "NO GO": colors.HexColor("#FFC7CE")}

    def table(rows, widths, recs=None):
        tb = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1)
        st = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
              ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]
        for i, r in enumerate(recs or [], 1):
            st.append(("BACKGROUND", (-2, i), (-2, i), rc.get(r, colors.white)))
        tb.setStyle(TableStyle(st))
        return tb

    def opp_rows(opps):
        hdr = ["#", "Project", "Client", "Ctry", "Fit", "Risk", "Win", "Fee $M", "Rec.", "Due"]
        rows = [hdr] + [[o.get("priority_rank"), P(o["project"]), P(o["client"]), o["country"], f"{o['fit']['score']:.0f}",
                         f"{o['risk']['score']:.0f}", f"{o['win']['probability']:.0%}", _m(o["fit"]["consulting_value_usd"]),
                         P(o["decision"]["recommendation"]), o.get("bid_due_date") or "TBC"] for o in opps]
        return table(rows, [8, 85, 55, 18, 10, 10, 10, 14, 24, 20], [o["decision"]["recommendation"] for o in opps])

    s = rd["stats"]
    E = [Paragraph("ECEC BD Intelligence: Daily Executive Report", h1),
         P(f"{t.isoformat()} (AST) | Model v{cfg['model_version']} | Pipeline (GO + caution): USD {_m(rd['pipe_fee'])}M fees, "
           f"USD {_m(rd['pipe_ev'])}M risk-weighted", body), Spacer(1, 4 * mm),
         Paragraph("Executive Summary", h2),
         P(f"{s.get('new', 0)} new, {s.get('updated', 0)} updated, {s.get('rejected', 0)} rejected (technical alignment). Active: "
           f"{rd['recs'].get('GO', 0)} GO / {rd['recs'].get('GO WITH CAUTION', 0)} GO WITH CAUTION / {rd['recs'].get('NO GO', 0)} NO GO. "
           f"{len(rd['upcoming'])} submissions due within 30 days. The fee is unknown on {rd['unknown_fee']} items, so their economics are provisional.", body),
         Paragraph("1. New Opportunities", h2), opp_rows(rd["new"]) if rd["new"] else P("None today.", body),
         PageBreak(), Paragraph("2. Top 20 Priority Opportunities", h2), opp_rows(rd["ranked"][:20]),
         Paragraph("Rationale: top 5", h2)]
    E += [P(f"<b>{o['project']}</b>: {'; '.join(o['decision']['reasons'])}. {o['win']['rationale']}".replace("&", "&amp;"), small) for o in rd["ranked"][:5]]
    E += [Paragraph("3. High-Risk Opportunities", h2)]
    E += [P(f"{o['project']}: risk {o['risk']['score']} ({o['risk']['classification']}); top risks: {', '.join(o['risk']['top_risks'])}", body) for o in rd["high_risk"]] \
         or [P("No active opportunity scores above 60.", body)]
    E += [Paragraph("4. Upcoming Deadlines (30 days)", h2),
          table([["Due", "Days", "Project", "Client", "Rec.", "Status"]] +
                [[o["bid_due_date"], _days(o, t), P(o["project"]), P(o["client"]), P(o["decision"]["recommendation"]), o["status"]] for o in rd["upcoming"]],
                [22, 12, 110, 60, 28, 22], [o["decision"]["recommendation"] for o in rd["upcoming"]]) if rd["upcoming"] else P("None.", body),
          Paragraph("5. Approved Opportunities", h2)]
    E += [P(f"{o['project']}: {o['status']}; capture plan {o.get('capture_plan')}", body) for o in rd["approved"]] or [P("None yet.", body)]
    E += [PageBreak(), Paragraph("6. Major Market Developments", h2),
          table([["Date", "Ctry", "Category", "Headline", "Implication"]] +
                [[m.get("date"), P(m.get("country")), P(m.get("category")), P(m.get("headline")), P(m.get("implication"))] for m in rd["recent_mi"][:15]],
                [18, 20, 28, 95, 105])]
    if trends:
        E += [Paragraph("Weekly Trend Analysis", h2),
              P(f"Intel by category: {dict(trends['mi_by_category'])} (vs prior week {trends['mi_delta']:+d}). New opportunities by sector: "
                f"{dict(trends['new_opps_by_sector'])}. Share not rated NO GO: {trends['go_rate']}.", body)]
    E += [Paragraph("7. Recommended Actions", h2),
          table([["Target", "#", "Action", "Impact", "Effort", "Owner"]] +
                [[a["target_date"], a["rank"], P(f"{a['action']} ({a['project'][:45]})"), P(a["expected_impact"]), a["effort"], P(a["responsible"])] for a in rd["actions"]],
                [20, 8, 130, 60, 14, 32]),
          Spacer(1, 4 * mm),
          P(f"Calibration: {cal['n_outcomes']} outcomes recorded; shift {cal['shift']:+.3f}. Full per-criterion justification is in opportunities.xlsx, sheet 'Scoring Audit'.", small)]
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm,
                      bottomMargin=10 * mm, title=f"ECEC BD Daily Report {t}").build(E)


def build_all(db, cfg, kb, today, run_stats, weekly):
    from . import learning
    mi = load_json(DATA / "market_intelligence.json", [])
    cal = learning.calibration(cfg)
    rd = report_data(db, cfg, today, mi, run_stats)
    trends = weekly_trends(db, mi, today) if weekly else None
    day = OUTPUT / "reports" / today.isoformat()
    xlsx = OUTPUT / "opportunities.xlsx"
    write_excel(db, cfg, today, mi, xlsx)
    md = markdown(rd, cfg, cal, trends)
    (day).mkdir(parents=True, exist_ok=True)
    (day / "daily_executive_report.md").write_text(md, encoding="utf-8")
    (OUTPUT / "LATEST_REPORT.md").write_text(md, encoding="utf-8")
    pdf(rd, cfg, cal, trends, day / "daily_executive_report.pdf")
    import shutil
    shutil.copy(xlsx, day / "opportunities.xlsx")
    return [str(p.relative_to(OUTPUT.parent)) for p in (xlsx, day / "daily_executive_report.md", day / "daily_executive_report.pdf")]
